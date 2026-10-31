#!/usr/bin/env python3
"""Generate the college System Design Project Report as .docx, then .pdf.

Layout follows the official format file (AMRITA School of Engineering,
23ECE302 assignment + 23ECE383 system design). All content comes from a JSON
data file; nothing is invented. Missing content renders as a bold red
"NOT YET AVAILABLE: <what is needed>" line, and missing metrics render "n/a".

Schema and rules: report/gpdk045/README_report.md
Example data:     report/gpdk045/report_data.example.json

Usage:
  /home/parshu/vm/rv/bin/python scripts/build_academic_docx.py \
      --data report/gpdk045/report_data.json --out /path/to/outdir [--no-pdf]
"""

import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

FONT = "Times New Roman"
BODY_PT = 12
HEADING_PT = 14
RED = RGBColor(0xC0, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
MAX_FIG_W_IN = 6.0
MAX_FIG_H_IN = 7.5
TEXT_W_IN = 6.5  # Letter width 8.5 in minus two 1 in margins
DOCX_NAME = "23ECE383_System_Design_Report.docx"

# Figure slots in the order they appear in the document. Figure numbers are
# fixed by this order so that proof references stay stable when slots fill in.
FIG_SLOTS = [
    ("circuit_diagram", "circuit diagram of the divider"),
    ("impl_screenshot", "simulator screenshot of the circuit implementation (without sizing)"),
    ("waveform_without_sizing", "simulator DC and transient waveform screenshots (without sizing)"),
    ("sized_circuit_screenshot", "simulator screenshot of the circuit based on the sizing concept"),
    ("waveform_with_sizing", "simulator DC and transient waveform screenshots (with sizing)"),
    ("pvt_screenshot", "PVT analysis screenshot"),
    ("layout_screenshot", "layout screenshot"),
    ("layout_waveform", "simulator waveform screenshot of the layout"),
]
FIG_NO = {key: i for i, (key, _) in enumerate(FIG_SLOTS, start=1)}
EXTRA_FIG_START = len(FIG_SLOTS) + 1

METRIC_UNITS = {
    "f_in_ghz": "GHz",
    "f_out_mhz": "MHz",
    "power_mw": "mW",
    "delay_ps": "ps",
    "pdp_fj": "fJ",
    "max_phase_step_err_deg": "deg",
    "swing_v": "V",
    "m7_vds_v": "V",
    "m7_vdsat_v": "V",
}

# (row label, metric key) for the values table; the saturation row is special.
VALUE_ROWS = [
    ("Power calculation (P)", "power_mw"),
    ("Delay calculation (t_pd)", "delay_ps"),
    ("Power Delay Product (PDP)", "pdp_fj"),
    ("Input frequency (f_in)", "f_in_ghz"),
    ("Output frequency (f_out)", "f_out_mhz"),
    ("Max phase-step error", "max_phase_step_err_deg"),
    ("Output swing", "swing_v"),
    ("Tail-device saturation check (M7)", None),
]


def na(value):
    return "n/a" if value is None else str(value)


def metric_text(metrics, key):
    value = metrics.get(key)
    if value is None:
        return "n/a"
    return f"{value} {METRIC_UNITS[key]}"


class Builder:
    def __init__(self, data, data_dir):
        self.d = data
        self.dir = data_dir
        self.doc = Document()
        self.tab_no = 0
        self.missing = []

    # ------------------------------------------------------------------ low level
    @staticmethod
    def _font(run, size=BODY_PT, bold=False, italic=False, color=None):
        run.font.name = FONT
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), FONT)
        rfonts.set(qn("w:cs"), FONT)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color is not None:
            run.font.color.rgb = color

    def text(self, p, s, size=BODY_PT, bold=False, italic=False, color=None):
        run = p.add_run(s)
        self._font(run, size, bold, italic, color)
        return run

    def para(self, s="", *, size=BODY_PT, bold=False, italic=False, color=None,
             align=None, before=0, after=6, keep=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep
        if align == "c":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "r":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if s:
            self.text(p, s, size, bold, italic, color)
        return p

    def heading(self, s):
        p = self.doc.add_paragraph(style="Heading 1")
        pf = p.paragraph_format
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        self.text(p, s, HEADING_PT, bold=True, color=BLACK)
        return p

    def need(self, what):
        """Record and render a visible NOT YET AVAILABLE line."""
        self.missing.append(what)
        self.para(f"NOT YET AVAILABLE: {what}", bold=True, color=RED, after=8)

    @staticmethod
    def _aspect(path):
        # PIL is used for the pixel size only. Explicit width and height are
        # passed to python-docx, which otherwise divides by the image DPI
        # (zero in some JPEGs).
        with Image.open(path) as im:
            w_px, h_px = im.size
        return h_px / w_px

    def _resolve(self, file_name):
        if not file_name:
            return None
        path = Path(file_name)
        return path if path.is_absolute() else (self.dir / path)

    def value_line(self, value, what):
        if value:
            self.para(str(value), after=8)
        else:
            self.need(what)

    def body(self, items, what):
        if isinstance(items, str):
            items = [items]
        if not items:
            self.need(what)
            return
        for item in items:
            if item:
                self.para(str(item), after=6)

    def _field(self, p, instr):
        def run_with(element=None, text=None):
            run = p.add_run(text) if text is not None else p.add_run()
            self._font(run)
            if element is not None:
                run._element.append(element)
            return run

        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        run_with(begin)
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = f" {instr} "
        run_with(instr_el)
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        run_with(sep)
        run_with(text="1")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run_with(end)

    def _cell(self, cell, s, bold=False, italic=False, color=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        self.text(p, s, bold=bold, italic=italic, color=color)

    def table(self, header, rows, widths, caption):
        self.tab_no += 1
        cap = self.para(keep=True, after=4)
        self.text(cap, f"Table {self.tab_no}: ", bold=True)
        self.text(cap, caption)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, h in enumerate(header):
            self._cell(t.rows[0].cells[i], h, bold=True)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        header_flag = OxmlElement("w:tblHeader")
        trpr.append(header_flag)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                if isinstance(value, tuple):
                    self._cell(cells[i], value[0], bold=True, color=value[1])
                else:
                    self._cell(cells[i], value)
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
            for row in t.rows:
                row.cells[i].width = Inches(w)
        self.para(after=4)
        return t

    def figure(self, n, slot, what):
        if not slot:
            self.need(f"Figure {n} - {what}")
            return
        path = self._resolve(slot.get("file"))
        if path is None or not path.is_file():
            self.need(f"Figure {n} - {what} (file not found: {slot.get('file')})")
            return
        caption = slot.get("caption")
        if not caption:
            self.need(f"Figure {n} caption text ({what})")
        aspect = self._aspect(path)  # height / width
        width_in = MAX_FIG_W_IN
        if aspect * width_in > MAX_FIG_H_IN:
            width_in = MAX_FIG_H_IN / aspect
        pic = self.para(align="c", keep=True, after=2)
        pic.add_run().add_picture(str(path), width=Inches(width_in),
                                  height=Inches(width_in * aspect))
        cap = self.para(align="c", after=10)
        self.text(cap, f"Figure {n}: ", bold=True)
        self.text(cap, caption or "")

    # --------------------------------------------------------------- page setup
    def _setup_styles(self):
        def fix_theme(style_el):
            rpr = style_el.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(attr)) is not None:
                    del rfonts.attrib[qn(attr)]
            rfonts.set(qn("w:eastAsia"), FONT)
            rfonts.set(qn("w:cs"), FONT)
            color = rpr.find(qn("w:color"))
            if color is not None:
                for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
                    if color.get(qn(attr)) is not None:
                        del color.attrib[qn(attr)]

        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(BODY_PT)
        fix_theme(normal.element)

        heading = self.doc.styles["Heading 1"]
        heading.font.name = FONT
        heading.font.size = Pt(HEADING_PT)
        heading.font.bold = True
        heading.font.italic = False
        heading.font.color.rgb = BLACK
        fix_theme(heading.element)

    def _page_setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        sec.footer_distance = Inches(0.5)

        # Double page border, as in the format file.
        sect_pr = sec._sectPr
        borders = OxmlElement("w:pgBorders")
        borders.set(qn("w:offsetFrom"), "text")
        borders.set(qn("w:display"), "allPages")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "double")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "10")
            el.set(qn("w:color"), "000000")
            borders.append(el)
        sect_pr.find(qn("w:pgMar")).addnext(borders)

        footer_p = sec.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.text(footer_p, "Page ")
        self._field(footer_p, "PAGE")
        self.text(footer_p, " of ")
        self._field(footer_p, "NUMPAGES")

    # ------------------------------------------------------------- title page
    def title_page(self):
        d = self.d
        logo = d.get("logo")
        if logo:
            path = self._resolve(logo)
            if path is not None and path.is_file():
                p = self.para(align="c", after=6)
                w_in = 3.6
                p.add_run().add_picture(str(path), width=Inches(w_in),
                                        height=Inches(w_in * self._aspect(path)))
            else:
                self.need(f"institution logo (file not found: {logo})")
        else:
            self.need('institution logo image (set "logo" in the data file)')

        self.para("AMRITA SCHOOL OF ENGINEERING, BENGALURU", align="c", after=10)
        self.para("DEPARTMENT OF ELECTRONICS & COMMUNICATION ENGINEERING", align="c", after=10)
        self.para("B. TECH. V SEMESTER (ECE AND EAC) AY 2026-2027", bold=True, align="c", before=10, after=10)
        self.para("ASSIGNMENT FOR 23ECE302 VLSI DESIGN", bold=True, align="c", before=10, after=10)
        self.para("SYSTEM DESIGN FOR 23ECE383 VLSI DESIGN LABORATORY", bold=True, align="c", before=10, after=6)
        self.para("REPORT", bold=True, align="c", after=4)
        self.para("JULY TO NOVEMBER 2026", bold=True, align="c", after=4)

        system_title = (d.get("titles") or {}).get("system_design_project")
        if system_title:
            self.para(system_title, bold=True, italic=True, align="c", before=14, after=14)
        else:
            self.need("system design project title (titles.system_design_project)")

        self.para("Submitted by", bold=True, align="c", before=8, after=6)
        sub = self.doc.add_table(rows=5, cols=3)
        sub.autofit = False
        headers = ["REGISTER NUMBER", "NAME", "SIGNATURE WITH DATE"]
        for i, h in enumerate(headers):
            self._cell(sub.rows[0].cells[i], h, italic=True)
        for i, w in enumerate([2.3, 2.0, 2.2]):
            sub.columns[i].width = Inches(w)
            for row in sub.rows:
                row.cells[i].width = Inches(w)

        self.para("BATCH NUMBER:", before=22, after=6)
        self.para("FACULTY EXAMINER SIGNATURE WITH EVALUATION DATE", bold=True, before=22, after=0)

    # --------------------------------------------------------------- sections
    def sec_title_and_paper(self):
        d = self.d
        titles = d.get("titles") or {}
        paper = d.get("paper") or {}

        self.heading("TITLE OF THE SYSTEM DESIGN PROJECT:")
        self.value_line(titles.get("system_design_project"),
                        "system design project title (titles.system_design_project)")

        self.heading("TITLE OF THE PAPER:")
        self.value_line(paper.get("title"), "paper title (paper.title)")

        self.heading("DETAILS OF THE PAPER (SOURCE OF THE TITLE):")
        if not paper:
            self.need("paper details (paper.*)")
        else:
            fields = [
                ("Authors", "authors"),
                ("Journal", "journal"),
                ("Volume", "volume"),
                ("Issue", "issue"),
                ("Pages", "pages"),
                ("Date", "date"),
                ("DOI", "doi"),
            ]
            for label, key in fields:
                p = self.para(after=2)
                self.text(p, f"{label}: ", bold=True)
                self.text(p, na(paper.get(key)))
            self.para(after=4)

    def sec_circuit(self):
        self.heading("CIRCUIT DIAGRAM:")
        self.figure(FIG_NO["circuit_diagram"], self._slot("circuit_diagram"),
                    dict(FIG_SLOTS)["circuit_diagram"])

    def sec_theory_tools(self):
        self.heading("THEORY AND WORKING PRINCIPLE:")
        self.body(self.d.get("theory"), "theory and working principle text (theory)")
        self.heading("TOOL & TECHNOLOGY USED:")
        self.body(self.d.get("tool_technology"), "tool and technology text (tool_technology)")

    def _slot(self, key):
        return (self.d.get("figures") or {}).get(key)

    def _design(self, name):
        return (self.d.get("designs") or {}).get(name) or {}

    def _device_table(self, name, caption):
        design = self._design(name)
        devices = design.get("devices")
        if not devices:
            self.need(f"device table for {name} (designs.{name}.devices)")
            return
        rows = [
            [dev.get("device"), na(dev.get("w")), na(dev.get("l")),
             na(dev.get("multiplier")), na(dev.get("role"))]
            for dev in devices
        ]
        self.table(["Device", "W", "L", "Multiplier", "Role"], rows,
                   [1.1, 1.2, 1.2, 1.1, 1.9], caption)

    def sec_implementation(self):
        self.heading("SCREENSHOT OF THE SIMULATOR BASED CIRCUIT IMPLEMENTATION:")
        design = self._design("without_sizing")
        self.body(design.get("description"),
                  "description of the circuit without sizing (designs.without_sizing.description)")
        self._device_table("without_sizing", "Device list, circuit without the sizing concept")
        self.figure(FIG_NO["impl_screenshot"], self._slot("impl_screenshot"),
                    dict(FIG_SLOTS)["impl_screenshot"])

        self.heading("SCREENSHOT OF THE OUTPUT WAVEFORM OBTAINED FROM THE SIMULATOR (DC AND TRANSIENT):")
        self.figure(FIG_NO["waveform_without_sizing"], self._slot("waveform_without_sizing"),
                    dict(FIG_SLOTS)["waveform_without_sizing"])

    def sec_sizing(self):
        self.heading("SCREENSHOT OF THE CIRCUIT BASED ON SIZING CONCEPT:")
        design = self._design("with_sizing")
        self.body(design.get("description"),
                  "description of the circuit with sizing (designs.with_sizing.description)")
        self._device_table("with_sizing", "Device list, circuit based on the sizing concept")
        self.figure(FIG_NO["sized_circuit_screenshot"], self._slot("sized_circuit_screenshot"),
                    dict(FIG_SLOTS)["sized_circuit_screenshot"])

        self.heading("SCREENSHOT OF THE OUTPUT WAVEFORM OBTAINED FROM THE SIMULATOR FOR THE CIRCUIT "
                     "WITH SIZING CONCEPT (DC AND TRANSIENT):")
        self.figure(FIG_NO["waveform_with_sizing"], self._slot("waveform_with_sizing"),
                    dict(FIG_SLOTS)["waveform_with_sizing"])

    def _sat_text(self, m):
        vds = m.get("m7_vds_v")
        vdsat = m.get("m7_vdsat_v")
        if vds is None or vdsat is None:
            return "n/a"
        saturated = vds >= vdsat
        text = f"Vds = {vds} V, Vdsat = {vdsat} V: " + ("saturated" if saturated else "NOT saturated")
        return text if saturated else (text, RED)

    def sec_values(self):
        self.heading("PROVIDE THE VALUES FOR THE FOLLOWING ALONG WITH THE SCREENSHOTS AS PROOF:")
        m_wo = self._design("without_sizing").get("metrics") or {}
        m_w = self._design("with_sizing").get("metrics") or {}
        proof = (f"Fig. {FIG_NO['waveform_without_sizing']} (without), "
                 f"Fig. {FIG_NO['waveform_with_sizing']} (with)")
        rows = []
        for label, key in VALUE_ROWS:
            if key is None:
                left, right = self._sat_text(m_wo), self._sat_text(m_w)
            else:
                left, right = metric_text(m_wo, key), metric_text(m_w, key)
            rows.append([label, left, right, proof])
        self.table(["Parameter", "Value without Sizing", "Value with Sizing", "Proof (figure)"],
                   rows, [1.9, 1.6, 1.6, 1.4],
                   "Performance values with and without the sizing concept, "
                   "with the proof figures for each row")

    def sec_pvt_layout(self):
        self.heading("SCREENSHOT OF THE PVT ANALYSIS:")
        pvt_rows = self.d.get("pvt_rows")
        if not pvt_rows:
            self.need("PVT corner sweep rows (pvt_rows)")
        else:
            rows = []
            for r in pvt_rows:
                divides = r.get("divides")
                if divides is True:
                    div_cell = "PASS"
                elif divides is False:
                    div_cell = ("FAIL", RED)
                else:
                    div_cell = "n/a"
                rows.append([
                    na(r.get("corner")), na(r.get("temp")), na(r.get("vdd")), div_cell,
                    na(r.get("f_out_mhz")), na(r.get("delay_ps")), na(r.get("power_mw")),
                    na(r.get("pdp_fj")), r.get("note") or "",
                ])
            self.table(["Corner", "Temp (C)", "VDD (V)", "Divides", "f_out (MHz)",
                        "Delay (ps)", "Power (mW)", "PDP (fJ)", "Note"],
                       rows, [0.75, 0.55, 0.55, 0.7, 0.75, 0.65, 0.7, 0.6, 1.25],
                       "PVT analysis results (FAIL rows are kept)")
        self.figure(FIG_NO["pvt_screenshot"], self._slot("pvt_screenshot"),
                    dict(FIG_SLOTS)["pvt_screenshot"])

        self.heading("SCREENSHOT OF THE LAYOUT:")
        self.figure(FIG_NO["layout_screenshot"], self._slot("layout_screenshot"),
                    dict(FIG_SLOTS)["layout_screenshot"])

        self.heading("SCREENSHOT OF THE OUTPUT WAVEFORM OBTAINED FROM THE SIMULATOR FOR THE LAYOUT:")
        self.figure(FIG_NO["layout_waveform"], self._slot("layout_waveform"),
                    dict(FIG_SLOTS)["layout_waveform"])

    def sec_rest(self):
        d = self.d
        self.heading("ANY OTHER RELEVANT INFORMATION:")
        self.body(d.get("other_info"), "other relevant information text (other_info)")
        extras = self._slot_list("extra_figures")
        for i, item in enumerate(extras):
            self.figure(EXTRA_FIG_START + i, item, f"additional figure {i + 1}")

        self.heading("CONCLUSION:")
        self.body(d.get("conclusion"), "conclusion text (conclusion)")

        self.heading("PROBLEMS FACED DURING THE IMPLEMENTATION:")
        self.body(d.get("problems_faced"), "problems faced text (problems_faced)")

        self.heading("PROBLEMS RESOLVED USING SOLUTIONS:")
        self.body(d.get("problems_resolved"), "problems resolved text (problems_resolved)")

        self.heading("ANY OTHER COMMENT:")
        self.body(d.get("other_comment"), "other comment text (other_comment)")

        self.heading("REFERENCES:")
        self.body(d.get("references"), "references (references)")

    def _slot_list(self, key):
        items = (self.d.get("figures") or {}).get(key)
        return items or []

    # ------------------------------------------------------------------ driver
    def build(self):
        self._setup_styles()
        self._page_setup()
        if self.d.get("example_only") is True:
            self.para("EXAMPLE DATA - NOT A RESULT", size=HEADING_PT, bold=True,
                      color=RED, align="c", after=12)
        self.title_page()
        self.doc.add_page_break()
        self.sec_title_and_paper()
        self.sec_circuit()
        self.sec_theory_tools()
        self.sec_implementation()
        self.sec_sizing()
        self.sec_values()
        self.sec_pvt_layout()
        self.sec_rest()

        core = self.doc.core_properties
        core.author = ""
        core.last_modified_by = ""
        core.title = (self.d.get("titles") or {}).get("system_design_project") or "System Design Project Report"


def convert_to_pdf(docx_path, out_dir):
    exe = shutil.which("libreoffice") or shutil.which("soffice")
    if exe is None:
        raise SystemExit("libreoffice not found on PATH")
    subprocess.run(
        [exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
        check=True, timeout=300, capture_output=True,
    )
    pdf_path = out_dir / (docx_path.stem + ".pdf")
    if not pdf_path.is_file():
        raise SystemExit(f"PDF was not produced: {pdf_path}")
    return pdf_path


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data", required=True, help="report data JSON (see README_report.md)")
    parser.add_argument("--out", required=True, help="output directory for the .docx and .pdf")
    parser.add_argument("--no-pdf", action="store_true", help="only write the .docx")
    args = parser.parse_args(argv)

    data_path = Path(args.data).resolve()
    data = json.loads(data_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise SystemExit("data file must contain a JSON object")

    builder = Builder(data, data_path.parent)
    builder.build()

    out_dir = Path(args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / DOCX_NAME
    builder.doc.save(str(docx_path))
    print(f"DOCX: {docx_path}")

    if not args.no_pdf:
        pdf_path = convert_to_pdf(docx_path, out_dir)
        print(f"PDF:  {pdf_path}")

    print(f"NOT YET AVAILABLE items: {len(builder.missing)}")
    for item in builder.missing:
        print(f"  - {item}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
